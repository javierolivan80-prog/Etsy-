"""Block 5a: real screenshots of the products, cropped and ready for the mockups."""
import os, subprocess, glob
import fitz
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(ROOT, "output", "imagenes", "_shots")
TMP = os.path.join(SHOTS, "_tmp")
os.makedirs(TMP, exist_ok=True)

def soffice(args):
    subprocess.run(["soffice", "--headless"] + args, check=True,
                   capture_output=True, timeout=240)

def autocrop(img, bg_tol=18):
    """Crop uniform background margins (cream/white)."""
    px = img.convert("RGB")
    w, h = px.size
    corner = px.getpixel((2, 2))
    def is_bg(p): return all(abs(p[i]-corner[i]) <= bg_tol for i in range(3))
    left, right, top, bot = w, 0, h, 0
    data = px.load()
    for y in range(0, h, 3):
        for x in range(0, w, 3):
            if not is_bg(data[x, y]):
                left = min(left, x); right = max(right, x)
                top = min(top, y); bot = max(bot, y)
    if right <= left: return img
    pad = 14
    return img.crop((max(0, left-pad), max(0, top-pad),
                     min(w, right+pad), min(h, bot+pad)))

def render_pdf_page(pdf, page, out, dpi=160, crop=True):
    d = fitz.open(pdf)
    pix = d[page].get_pixmap(dpi=dpi)
    pix.save(out)
    if crop:
        img = Image.open(out)
        autocrop(img).save(out)
    return out

# ---------- Product 1: sample-data planner -> per-sheet screenshots ----------
def product1():
    sample = os.path.join(TMP, "sample-planner.xlsx")
    subprocess.run(["python3", os.path.join(ROOT, "scripts", "block1_budget.py"),
                    "sample", sample], check=True)
    # landscape + fit-to-width for clean single-page-per-sheet export
    from openpyxl import load_workbook
    from openpyxl.worksheet.properties import PageSetupProperties
    wb = load_workbook(sample)
    for ws in wb:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.save(sample)
    soffice(["--convert-to", "pdf", "--outdir", TMP, sample])
    pdf = os.path.join(TMP, "sample-planner.pdf")
    d = fitz.open(pdf)
    # map sheet -> first page containing its title text
    marks = {"p1-dashboard": "2026 BUDGET PLANNER", "p1-income": "INCOME SOURCE",
             "p1-expenses": "CATEGORY", "p1-savings": "SAVINGS GOAL",
             "p1-debts": "STARTING BALANCE", "p1-annual": "ANNUAL SUMMARY"}
    used = set()
    for key, needle in marks.items():
        for i in range(d.page_count):
            if i in used: continue
            if needle in d[i].get_text():
                render_pdf_page(pdf, i, os.path.join(SHOTS, key + ".png"))
                used.add(i); break
        else:
            raise AssertionError(f"sheet page not found for {key}")
    print("product1 shots:", sorted(marks))

# ---------- Product 2: planner PDF pages ----------
def product2():
    pdf = os.path.join(ROOT, "output", "producto-2", "PlanaryShop-Weekly-Planner-US-Letter.pdf")
    pages = {"p2-cover": 0, "p2-howto": 1, "p2-goals": 2, "p2-action": 3,
             "p2-monthly": 4, "p2-weekly": 6, "p2-habits": 12, "p2-notes": 14,
             "p2-dotted": 16}
    for key, pg in pages.items():
        render_pdf_page(pdf, pg, os.path.join(SHOTS, key + ".png"), crop=False)
    print("product2 shots:", sorted(pages))

# ---------- Product 3: docx + tracker -> screenshots ----------
def product3():
    base = os.path.join(ROOT, "output", "producto-3")
    docs = {"p3-invoice": "PlanaryShop-Invoice-Template.docx",
            "p3-quote": "PlanaryShop-Quote-Template.docx",
            "p3-agreement": "PlanaryShop-Service-Agreement-Template.docx",
            "p3-cancel": "PlanaryShop-Cancellation-Policy-Template.docx"}
    for key, f in docs.items():
        soffice(["--convert-to", "pdf", "--outdir", TMP, os.path.join(base, f)])
        pdf = os.path.join(TMP, f.replace(".docx", ".pdf"))
        render_pdf_page(pdf, 0, os.path.join(SHOTS, key + ".png"), crop=False)
    # tracker: landscape fit
    from openpyxl import load_workbook
    from openpyxl.worksheet.properties import PageSetupProperties
    src = os.path.join(base, "PlanaryShop-Client-Tracker.xlsx")
    tmpx = os.path.join(TMP, "tracker-sample.xlsx")
    wb = load_workbook(src)
    ws = wb["Clients"]
    sample_rows = [
        ("Ava Thompson","ava@email.com","555-0142","Logo design","In progress",450,450,225,"Mar 14","Loves option B"),
        ("Miguel Reyes","miguel@email.com","555-0177","House cleaning","Booked",120,120,120,"Mar 20","Biweekly slot"),
        ("Dana Lee","dana@email.com","555-0119","Family photos","Quoted",350,0,0,"Mar 12","Waiting on date"),
        ("Sam Carter","sam@email.com","555-0163","Coaching pack","Paid",600,600,600,"Apr 02","Renewal likely"),
    ]
    for i, row in enumerate(sample_rows):
        r = 7 + i
        vals = [row[0], row[1], row[2], row[3], row[4], row[5], row[6], row[7], None, row[8], row[9]]
        for j, v in enumerate(vals):
            if v is not None:
                ws.cell(row=r, column=2+j, value=v)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.save(tmpx)
    soffice(["--convert-to", "pdf", "--outdir", TMP, tmpx])
    render_pdf_page(os.path.join(TMP, "tracker-sample.pdf"), 0,
                    os.path.join(SHOTS, "p3-tracker.png"))
    print("product3 shots: 4 docs + tracker")

if __name__ == "__main__":
    product1(); product2(); product3()
    for f in sorted(glob.glob(os.path.join(SHOTS, "*.png"))):
        im = Image.open(f)
        print(os.path.basename(f), im.size)
