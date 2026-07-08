"""Block 6: final verification — every deliverable listed with size + hard checks."""
import os, re, sys, zipfile, html
import fitz
from PIL import Image
from openpyxl import load_workbook
from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "output")
ok_all = True
report = []

def check(label, cond, detail=""):
    global ok_all
    mark = "✔" if cond else "✘"
    if not cond: ok_all = False
    report.append(f"  [{mark}] {label}" + (f" — {detail}" if detail else ""))
    return cond

def size(p):
    kb = os.path.getsize(p) / 1024
    return f"{kb:,.0f} KB"

def xlsx_check(path, min_formulas):
    wb = load_workbook(path)
    n = sum(1 for ws in wb for row in ws.iter_rows()
            for c in row if isinstance(c.value, str) and c.value.startswith("="))
    return check(os.path.basename(path) + f" ({size(path)})", n >= min_formulas,
                 f"abre OK, {n} fórmulas")

def pdf_check(path, pages):
    d = fitz.open(path)
    return check(os.path.basename(path) + f" ({size(path)})", d.page_count == pages,
                 f"{d.page_count} páginas (esperadas {pages})")

def docx_check(path):
    d = Document(path)
    texts = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            for cell in row.cells: texts.append(cell.text)
    full = "\n".join(texts)
    bad = any(t in full for t in ["TODO", "XXX", "Lorem", "{{", "}}"])
    balanced = full.count("[") == full.count("]")
    fields = len(re.findall(r"\[[^\]]+\]", full))
    return check(os.path.basename(path) + f" ({size(path)})",
                 not bad and balanced and fields >= 5,
                 f"{fields} campos [..], sin placeholders rotos")

print("=" * 64)
print("  PLANARYSHOP — VERIFICACIÓN FINAL")
print("=" * 64)

report.append("\nPRODUCTO 1 — Budget Planner 2026")
p1 = os.path.join(OUT, "producto-1")
xlsx_check(os.path.join(p1, "PlanaryShop-Budget-Planner-2026.xlsx"), 100)
xlsx_check(os.path.join(p1, "PlanaryShop-Budget-Planner-2026-ES.xlsx"), 100)
pdf_check(os.path.join(p1, "Quick-Start-Guide.pdf"), 1)
pdf_check(os.path.join(p1, "Quick-Start-Guide-ES.pdf"), 1)

report.append("\nPRODUCTO 2 — Weekly Planner + Habit Tracker")
p2 = os.path.join(OUT, "producto-2")
pdf_check(os.path.join(p2, "PlanaryShop-Weekly-Planner-US-Letter.pdf"), 17)
pdf_check(os.path.join(p2, "PlanaryShop-Weekly-Planner-A4.pdf"), 17)

report.append("\nPRODUCTO 3 — Small Business Starter Kit")
p3 = os.path.join(OUT, "producto-3")
for f in ["PlanaryShop-Invoice-Template.docx", "PlanaryShop-Quote-Template.docx",
          "PlanaryShop-Service-Agreement-Template.docx",
          "PlanaryShop-Cancellation-Policy-Template.docx"]:
    docx_check(os.path.join(p3, f))
xlsx_check(os.path.join(p3, "PlanaryShop-Client-Tracker.xlsx"), 30)
zp = os.path.join(p3, "PlanaryShop-Business-Kit.zip")
z = zipfile.ZipFile(zp)
check(f"PlanaryShop-Business-Kit.zip ({size(zp)})", len(z.namelist()) == 5,
      f"{len(z.namelist())} archivos dentro")

report.append("\nLISTINGS")
for f in sorted(os.listdir(os.path.join(OUT, "listings"))):
    path = os.path.join(OUT, "listings", f)
    txt = open(path, encoding="utf-8").read()
    m = re.search(r"=== TITLE \((\d+)/140", txt)
    tlen = int(m.group(1))
    tags = re.findall(r"^\s*\d+\. (.+?)  \(\d+\)$", txt, re.M)
    tags_ok = len(tags) == 13 and all(len(t) <= 20 for t in tags)
    has_parts = all(k in txt for k in ["WHAT'S INCLUDED", "HOW IT WORKS", "FAQ",
                                       "DIGITAL PRODUCT", "personal use"])
    check(f + f" ({size(path)})", tlen <= 140 and tags_ok and has_parts,
          f"título {tlen}/140, {len(tags)} tags ≤20, secciones completas")

report.append("\nIMÁGENES (2700×2025)")
n = 0
for sub in ["producto-1", "producto-2", "producto-3"]:
    d = os.path.join(OUT, "imagenes", sub)
    files = sorted(f for f in os.listdir(d) if f.endswith(".jpg"))
    sizes_ok = all(Image.open(os.path.join(d, f)).size == (2700, 2025) for f in files)
    kb = sum(os.path.getsize(os.path.join(d, f)) for f in files) / 1024
    n += len(files)
    check(f"imagenes/{sub} ({kb:,.0f} KB)", len(files) == 10 and sizes_ok,
          f"{len(files)}/10 JPG, todas 2700×2025")
check("total imágenes", n == 30, f"{n}/30")

print("\n".join(report))
print()
print("=" * 64)
if ok_all:
    print("  RESULTADO: TODO CORRECTO ✔  — listo para subir a Etsy")
else:
    print("  RESULTADO: HAY FALLOS ✘ — revisar arriba")
    sys.exit(1)
