"""PlanaryShop — Block 1: Budget Planner 2026 (EN/ES) + Quick-Start Guide PDF."""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.series import DataPoint

CREAM = "FAF7F2"; BEIGE = "E9DFD2"; SAGE = "8A9B85"; TERRA = "C98A6B"
INK = "2E2A26"; WGRAY = "4A443E"; WHITE = "FFFFFF"
CH_INC = "4E8A35"; CH_EXP = "C05F35"; CH_SAV = "00958A"  # validated chart palette
PIE15 = ["567455","C05F35","00958A","B5830A","5E7FB8","B07AA8","8A9B85","C98A6B",
         "3E6B78","9C4A4A","6B5B3E","7FA05A","D9A441","8A7BA8","4A443E"]
SERIF = "Georgia"; SANS = "Lato"
CUR = '$#,##0.00'; PCT = '0.0%'
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "output", "producto-1")

L = {
 "en": dict(
   months=["JAN","FEB","MAR","APR","MAY","JUN","JUL","AUG","SEP","OCT","NOV","DEC"],
   dashboard="Dashboard", income="Income", expenses="Expenses", savings="Savings",
   debts="Debts", summary="Annual Summary",
   title="2026 BUDGET PLANNER", sub_dash="Your whole year at a glance — every number updates automatically",
   sub_inc="Type your income in the white cells — beige cells calculate themselves",
   sub_exp="Track spending by category — beige cells calculate themselves",
   sub_sav="Set a target, log what you saved — progress fills in on its own",
   sub_deb="List each debt once — the remaining balance updates as you pay",
   sub_sum="Nothing to type here — this page reads from Income and Expenses",
   source="INCOME SOURCE", category="CATEGORY", total="TOTAL", month="MONTH",
   goal="SAVINGS GOAL", target="TARGET", saved="SAVED", progress="PROGRESS",
   debt="DEBT", start_bal="STARTING BALANCE", paid="PAID SO FAR", remaining="REMAINING",
   s_income="INCOME", s_expenses="EXPENSES", s_net="NET SAVINGS", s_rate="SAVINGS RATE",
   c_inc="TOTAL INCOME", c_exp="TOTAL EXPENSES", c_net="NET SAVED", c_rate="AVG SAVINGS RATE",
   chart_bar="Income, expenses & savings by month", chart_pie="Where the year's money went",
   sources=["Salary","Side hustle","Freelance","Investments","Other"],
   cats=["Housing","Utilities","Groceries","Transportation","Insurance","Health",
         "Dining Out","Entertainment","Subscriptions","Shopping","Personal Care",
         "Education","Gifts","Travel","Other"],
   goals=["Emergency fund","Vacation","New car","Home","Other"],
   debtnames=["Credit card","Student loan","Car loan","Medical","Other"],
   thanks="Thank you for supporting PlanaryShop! For personal use only.",
 ),
 "es": dict(
   months=["ENE","FEB","MAR","ABR","MAY","JUN","JUL","AGO","SEP","OCT","NOV","DIC"],
   dashboard="Panel", income="Ingresos", expenses="Gastos", savings="Ahorros",
   debts="Deudas", summary="Resumen Anual",
   title="PLANIFICADOR DE PRESUPUESTO 2026", sub_dash="Todo tu año de un vistazo — cada número se actualiza solo",
   sub_inc="Escribe tus ingresos en las celdas blancas — las beige se calculan solas",
   sub_exp="Registra tus gastos por categoría — las celdas beige se calculan solas",
   sub_sav="Define una meta, anota lo ahorrado — el progreso se rellena solo",
   sub_deb="Anota cada deuda una vez — el saldo pendiente se actualiza al pagar",
   sub_sum="Aquí no se escribe nada — esta hoja lee de Ingresos y Gastos",
   source="FUENTE DE INGRESO", category="CATEGORÍA", total="TOTAL", month="MES",
   goal="META DE AHORRO", target="OBJETIVO", saved="AHORRADO", progress="PROGRESO",
   debt="DEUDA", start_bal="SALDO INICIAL", paid="PAGADO", remaining="PENDIENTE",
   s_income="INGRESOS", s_expenses="GASTOS", s_net="AHORRO NETO", s_rate="TASA DE AHORRO",
   c_inc="INGRESOS TOTALES", c_exp="GASTOS TOTALES", c_net="AHORRO NETO", c_rate="TASA MEDIA DE AHORRO",
   chart_bar="Ingresos, gastos y ahorro por mes", chart_pie="A dónde fue el dinero del año",
   sources=["Salario","Ingreso extra","Freelance","Inversiones","Otro"],
   cats=["Vivienda","Suministros","Supermercado","Transporte","Seguros","Salud",
         "Restaurantes","Ocio","Suscripciones","Compras","Cuidado personal",
         "Educación","Regalos","Viajes","Otros"],
   goals=["Fondo de emergencia","Vacaciones","Coche nuevo","Hogar","Otro"],
   debtnames=["Tarjeta de crédito","Préstamo estudios","Préstamo coche","Médico","Otro"],
   thanks="¡Gracias por apoyar a PlanaryShop! Solo para uso personal.",
 ),
}

thin = Side(style="thin", color=BEIGE)
def fl(c): return PatternFill("solid", start_color=c)

def sheet(wb, name, tab=SAGE):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    ws.column_dimensions['A'].width = 2
    return ws

def paint_bg(ws, rows, cols):
    for r in range(1, rows+1):
        for c in range(1, cols+1):
            ws.cell(row=r, column=c).fill = fl(CREAM)

def header(ws, title, sub, cmax):
    t = ws.cell(row=2, column=2, value=title)
    t.font = Font(name=SERIF, size=20, bold=True, color=INK)
    s = ws.cell(row=3, column=2, value=sub)
    s.font = Font(name=SANS, size=10, color=WGRAY)
    for c in range(2, cmax+1):
        ws.cell(row=4, column=c).fill = fl(SAGE)
    ws.row_dimensions[2].height = 28; ws.row_dimensions[4].height = 3
    ws.row_dimensions[5].height = 8

def th(ws, r, c, text, align="left"):
    cell = ws.cell(row=r, column=c, value=text)
    cell.font = Font(name=SANS, size=9, bold=True, color=WHITE)
    cell.fill = fl(SAGE)
    cell.alignment = Alignment(horizontal=align, vertical="center",
                               indent=1 if align == "left" else 0)
    ws.row_dimensions[r].height = 20

def inp(ws, r, c, val=None, fmt=CUR, align="right"):
    cell = ws.cell(row=r, column=c, value=val)
    cell.font = Font(name=SANS, size=10, color=INK)
    cell.fill = fl(WHITE)
    cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    if fmt: cell.number_format = fmt
    cell.alignment = Alignment(horizontal=align, vertical="center",
                               indent=1 if align == "left" else 0)
    return cell

def fml(ws, r, c, formula, fmt=CUR, bold=False, align="right"):
    cell = ws.cell(row=r, column=c, value=formula)
    cell.font = Font(name=SANS, size=10, bold=bold, color=INK)
    cell.fill = fl(BEIGE)
    cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    if fmt: cell.number_format = fmt
    cell.alignment = Alignment(horizontal=align, vertical="center")
    return cell

def label(ws, r, c, text, bold=False):
    cell = ws.cell(row=r, column=c, value=text)
    cell.font = Font(name=SANS, size=10, bold=bold, color=INK)
    cell.fill = fl(WHITE)
    cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    return cell

def thanks_row(ws, r, cmax, text):
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=cmax)
    cell = ws.cell(row=r, column=2, value=text)
    cell.font = Font(name=SERIF, size=9, italic=True, color=TERRA)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[r].height = 22

def grid_sheet(wb, s, name, sub, rowhdr, rows, tab, sample_rows=None):
    """Month grid: rows x 12 months + TOTAL col + TOTAL row."""
    ws = sheet(wb, name, tab)
    paint_bg(ws, 20 + len(rows), 16)
    header(ws, name.upper(), sub, 16)
    ws.column_dimensions['B'].width = 20
    for j in range(12): ws.column_dimensions[get_column_letter(3+j)].width = 10.5
    ws.column_dimensions['O'].width = 12
    th(ws, 7, 2, rowhdr)
    for j, m in enumerate(s["months"]): th(ws, 7, 3+j, m, "right")
    th(ws, 7, 15, s["total"], "right")
    r0 = 8
    for i, nm in enumerate(rows):
        r = r0 + i
        label(ws, r, 2, nm)
        for j in range(12):
            v = None
            if sample_rows is not None: v = sample_rows[i][j]
            inp(ws, r, 3+j, v)
        fml(ws, r, 15, f"=SUM(C{r}:N{r})", bold=True)
        ws.row_dimensions[r].height = 19
    rt = r0 + len(rows)
    lab = ws.cell(row=rt, column=2, value=s["total"])
    lab.font = Font(name=SANS, size=10, bold=True, color=WHITE); lab.fill = fl(WGRAY)
    lab.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for j in range(12):
        cl = get_column_letter(3+j)
        cell = fml(ws, rt, 3+j, f"=SUM({cl}{r0}:{cl}{rt-1})", bold=True)
        cell.fill = fl(BEIGE)
    fml(ws, rt, 15, f"=SUM(O{r0}:O{rt-1})", bold=True)
    ws.row_dimensions[rt].height = 20
    thanks_row(ws, rt + 2, 15, s["thanks"])
    ws.freeze_panes = "C8"
    return ws, r0, rt

def build_planner(lang, path, sample=False):
    s = L[lang]
    wb = Workbook(); wb.remove(wb.active)

    sample_inc = sample_exp = None
    if sample:
        base_inc = [3200, 650, 400, 120, 0]
        sample_inc = [[round(b * (1 + 0.03*((i+j) % 3)), 0) if b else None
                       for j in range(12)] for i, b in enumerate(base_inc)]
        base_exp = [1400, 180, 520, 210, 160, 90, 140, 80, 45, 120, 60, 50, 40, 100, 35]
        sample_exp = [[round(b * (1 + 0.05*((i*j) % 2)), 0) for j in range(12)]
                      for i, b in enumerate(base_exp)]

    inc_ws, inc_r0, inc_rt = grid_sheet(wb, s, s["income"], s["sub_inc"],
                                        s["source"], s["sources"], SAGE, sample_inc)
    exp_ws, exp_r0, exp_rt = grid_sheet(wb, s, s["expenses"], s["sub_exp"],
                                        s["category"], s["cats"], TERRA, sample_exp)

    # ---- Savings ----
    sv = sheet(wb, s["savings"], SAGE)
    paint_bg(sv, 20, 8)
    header(sv, s["savings"].upper(), s["sub_sav"], 6)
    for col, w in zip("BCDE", [24, 14, 14, 14]): sv.column_dimensions[col].width = w
    th(sv, 7, 2, s["goal"]); th(sv, 7, 3, s["target"], "right")
    th(sv, 7, 4, s["saved"], "right"); th(sv, 7, 5, s["progress"], "right")
    sv_samples = [(8000, 3400), (2500, 900), (12000, 2100), (6000, 750), (None, None)]
    for i, g in enumerate(s["goals"]):
        r = 8 + i
        label(sv, r, 2, g)
        tgt, svd = sv_samples[i] if sample else (None, None)
        inp(sv, r, 3, tgt); inp(sv, r, 4, svd)
        fml(sv, r, 5, f'=IF(C{r}=0,0,D{r}/C{r})', PCT)
        sv.row_dimensions[r].height = 19
    thanks_row(sv, 15, 5, s["thanks"])

    # ---- Debts ----
    db = sheet(wb, s["debts"], TERRA)
    paint_bg(db, 20, 8)
    header(db, s["debts"].upper(), s["sub_deb"], 6)
    for col, w in zip("BCDE", [24, 16, 14, 14]): db.column_dimensions[col].width = w
    th(db, 7, 2, s["debt"]); th(db, 7, 3, s["start_bal"], "right")
    th(db, 7, 4, s["paid"], "right"); th(db, 7, 5, s["remaining"], "right")
    db_samples = [(4200, 1650), (18500, 3200), (9800, 4100), (750, 750), (None, None)]
    for i, d in enumerate(s["debtnames"]):
        r = 8 + i
        label(db, r, 2, d)
        sb, pd = db_samples[i] if sample else (None, None)
        inp(db, r, 3, sb); inp(db, r, 4, pd)
        fml(db, r, 5, f"=C{r}-D{r}")
        db.row_dimensions[r].height = 19
    thanks_row(db, 15, 5, s["thanks"])

    # ---- Annual Summary ----
    q = lambda n: f"'{n}'" if " " in n else n
    IN, EX = q(s["income"]), q(s["expenses"])
    an = sheet(wb, s["summary"], WGRAY)
    paint_bg(an, 26, 8)
    header(an, s["summary"].upper(), s["sub_sum"], 6)
    for col, w in zip("BCDEF", [12, 15, 15, 15, 15]): an.column_dimensions[col].width = w
    th(an, 7, 2, s["month"]); th(an, 7, 3, s["s_income"], "right")
    th(an, 7, 4, s["s_expenses"], "right"); th(an, 7, 5, s["s_net"], "right")
    th(an, 7, 6, s["s_rate"], "right")
    for j in range(12):
        r = 8 + j
        cl = get_column_letter(3+j)
        mc = ws_month = an.cell(row=r, column=2, value=s["months"][j])
        mc.font = Font(name=SANS, size=10, bold=True, color=WGRAY)
        mc.fill = fl(WHITE); mc.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        mc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        fml(an, r, 3, f"={IN}!{cl}{inc_rt}")
        fml(an, r, 4, f"={EX}!{cl}{exp_rt}")
        fml(an, r, 5, f"=C{r}-D{r}", bold=True)
        fml(an, r, 6, f"=IF(C{r}=0,0,E{r}/C{r})", PCT)
        an.row_dimensions[r].height = 19
    rt = 20
    lab = an.cell(row=rt, column=2, value=s["total"])
    lab.font = Font(name=SANS, size=10, bold=True, color=WHITE); lab.fill = fl(WGRAY)
    lab.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    fml(an, rt, 3, "=SUM(C8:C19)", bold=True); fml(an, rt, 4, "=SUM(D8:D19)", bold=True)
    fml(an, rt, 5, "=SUM(E8:E19)", bold=True)
    fml(an, rt, 6, "=IF(C20=0,0,E20/C20)", PCT, bold=True)
    thanks_row(an, 22, 6, s["thanks"])

    # ---- Dashboard ----
    AN = q(s["summary"])
    dash = sheet(wb, s["dashboard"], INK)
    paint_bg(dash, 46, 16)
    header(dash, s["title"], s["sub_dash"], 15)
    for col, w in zip("BCDEFGHIJKLMN", [14, 6, 14, 6, 14, 6, 14, 6, 10, 10, 10, 10, 10]):
        dash.column_dimensions[col].width = w
    cards = [(s["c_inc"], f"={AN}!C20", CUR, SAGE),
             (s["c_exp"], f"={AN}!D20", CUR, TERRA),
             (s["c_net"], f"={AN}!E20", CUR, WGRAY),
             (s["c_rate"], f"={AN}!F20", PCT, SAGE)]
    for i, (t, f, fmt, accent) in enumerate(cards):
        c0 = 2 + i*2
        for rr in range(7, 10):
            for cc in (c0, c0+1):
                cell = dash.cell(row=rr, column=cc)
                cell.fill = fl(WHITE)
                cell.border = Border(
                    left=Side(style="thick", color=accent) if cc == c0 else None,
                    top=thin if rr == 7 else None, bottom=thin if rr == 9 else None,
                    right=thin if cc == c0+1 else None)
        lab = dash.cell(row=7, column=c0, value=t)
        lab.font = Font(name=SANS, size=8, bold=True, color=WGRAY)
        lab.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        dash.merge_cells(start_row=8, start_column=c0, end_row=9, end_column=c0+1)
        v = dash.cell(row=8, column=c0, value=f)
        v.font = Font(name=SERIF, size=15, bold=True, color=INK)
        v.number_format = fmt
        v.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    dash.row_dimensions[7].height = 16
    dash.row_dimensions[8].height = 16; dash.row_dimensions[9].height = 16

    bar = BarChart(); bar.type = "col"; bar.grouping = "clustered"
    bar.title = s["chart_bar"]
    data = Reference(an, min_col=3, max_col=5, min_row=7, max_row=19)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(Reference(an, min_col=2, min_row=8, max_row=19))
    for ser, col in zip(bar.series, (CH_INC, CH_EXP, CH_SAV)):
        ser.graphicalProperties.solidFill = col
        ser.graphicalProperties.line.solidFill = WHITE
    bar.gapWidth = 60; bar.legend.position = 'b'
    bar.y_axis.numFmt = '$#,##0'
    bar.width = 24; bar.height = 9
    dash.add_chart(bar, "B12")

    pie = DoughnutChart(holeSize=58)
    pie.title = s["chart_pie"]
    pie.add_data(Reference(exp_ws, min_col=15, min_row=exp_r0, max_row=exp_rt-1),
                 titles_from_data=False)
    pie.set_categories(Reference(exp_ws, min_col=2, min_row=exp_r0, max_row=exp_rt-1))
    pts = []
    for i, col in enumerate(PIE15):
        dp = DataPoint(idx=i)
        dp.graphicalProperties.solidFill = col
        dp.graphicalProperties.line.solidFill = WHITE
        dp.graphicalProperties.line.width = 25400
        pts.append(dp)
    pie.series[0].data_points = pts
    pie.legend.position = 'r'
    pie.width = 24; pie.height = 10.5
    dash.add_chart(pie, "B31")
    thanks_row(dash, 45, 15, s["thanks"])

    order = [s["dashboard"], s["income"], s["expenses"], s["savings"], s["debts"], s["summary"]]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(path)
    return path

# ---------------- Quick-Start Guide PDF ----------------
GUIDE = {
 "en": dict(
   title="Quick-Start Guide", product="2026 Budget Planner",
   steps=[
     ("Open the planner in Excel or Google Sheets",
      "Works in Microsoft Excel (2016 or newer) and Google Sheets. In Sheets: File > Import > Upload."),
     ("Add your income on the Income tab",
      "Type each source of income in the white cells, one column per month. Totals fill in by themselves."),
     ("Log spending on the Expenses tab",
      "15 ready-made categories. Enter what you spend each month — the beige cells do the math."),
     ("Set goals in Savings and Debts",
      "Give each goal a target and log what you put aside. Debts show the remaining balance as you pay."),
     ("Watch the Dashboard come to life",
      "Charts and summary cards update automatically. Check Annual Summary for your monthly savings rate."),
   ],
   rule_h="The golden rule",
   rule="Only type in the WHITE cells. The beige cells contain formulas — if you overwrite a TOTAL cell, that number stops updating.",
   support="Questions? Message PlanaryShop on Etsy — we reply within 24 hours.",
   footer="Thank you for supporting PlanaryShop! For personal use only."),
 "es": dict(
   title="Guía Rápida", product="Planificador de Presupuesto 2026",
   steps=[
     ("Abre el planificador en Excel o Google Sheets",
      "Funciona en Microsoft Excel (2016 o posterior) y Google Sheets. En Sheets: Archivo > Importar > Subir."),
     ("Añade tus ingresos en la hoja Ingresos",
      "Escribe cada fuente de ingreso en las celdas blancas, una columna por mes. Los totales se rellenan solos."),
     ("Registra tus gastos en la hoja Gastos",
      "15 categorías listas para usar. Anota lo que gastas cada mes — las celdas beige hacen las cuentas."),
     ("Define metas en Ahorros y Deudas",
      "Pon un objetivo a cada meta y anota lo que apartas. Las deudas muestran el saldo pendiente al pagar."),
     ("Mira cómo cobra vida el Panel",
      "Los gráficos y tarjetas se actualizan solos. Consulta el Resumen Anual para ver tu tasa de ahorro mensual."),
   ],
   rule_h="La regla de oro",
   rule="Escribe solo en las celdas BLANCAS. Las celdas beige contienen fórmulas — si sobrescribes una celda TOTAL, ese número deja de actualizarse.",
   support="¿Dudas? Escribe a PlanaryShop en Etsy — respondemos en menos de 24 horas.",
   footer="¡Gracias por apoyar a PlanaryShop! Solo para uso personal."),
}

def build_guide(lang, path):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.colors import HexColor
    from reportlab.pdfgen import canvas
    cream = HexColor("#FAF7F2"); beige = HexColor("#E9DFD2"); sage = HexColor("#8A9B85")
    terra = HexColor("#C98A6B"); ink = HexColor("#2E2A26"); wgray = HexColor("#4A443E")
    g = GUIDE[lang]
    W, H = letter
    c = canvas.Canvas(path, pagesize=letter)
    c.setFillColor(cream); c.rect(0, 0, W, H, stroke=0, fill=1)
    c.setFillColor(sage); c.rect(0, H-8, W, 8, stroke=0, fill=1)
    c.setFillColor(ink); c.setFont("Times-Bold", 26)
    c.drawString(54, H-70, g["title"])
    c.setFillColor(terra); c.setFont("Helvetica-Bold", 11)
    c.drawString(54, H-90, ("PLANARYSHOP · " + g["product"]).upper())
    c.setStrokeColor(beige); c.setLineWidth(1.2); c.line(54, H-104, W-54, H-104)
    y = H - 140
    for i, (h, body) in enumerate(g["steps"], 1):
        c.setFillColor(sage); c.circle(70, y+4, 13, stroke=0, fill=1)
        c.setFillColor(HexColor("#FFFFFF")); c.setFont("Helvetica-Bold", 13)
        c.drawCentredString(70, y-0.5, str(i))
        c.setFillColor(ink); c.setFont("Times-Bold", 13.5)
        c.drawString(94, y+2, h)
        c.setFillColor(wgray); c.setFont("Helvetica", 9.5)
        # naive wrap at ~92 chars
        words = body.split(); line = ""; yy = y - 14
        for w_ in words:
            if len(line + " " + w_) > 92:
                c.drawString(94, yy, line.strip()); yy -= 12; line = w_
            else:
                line += " " + w_
        c.drawString(94, yy, line.strip())
        y = yy - 26
    # golden rule box
    bx_h = 64
    c.setFillColor(beige); c.roundRect(54, y-bx_h, W-108, bx_h, 8, stroke=0, fill=1)
    c.setFillColor(terra); c.setFont("Times-Bold", 12.5)
    c.drawString(70, y-22, "★  " + g["rule_h"])
    c.setFillColor(ink); c.setFont("Helvetica", 9.5)
    words = g["rule"].split(); line = ""; yy = y - 38
    for w_ in words:
        if len(line + " " + w_) > 100:
            c.drawString(70, yy, line.strip()); yy -= 12; line = w_
        else:
            line += " " + w_
    c.drawString(70, yy, line.strip())
    y = y - bx_h - 30
    c.setFillColor(sage); c.setFont("Helvetica-Bold", 10)
    c.drawString(54, y, g["support"])
    c.setFillColor(terra); c.setFont("Times-Italic", 9.5)
    c.drawCentredString(W/2, 40, g["footer"])
    c.setFillColor(sage); c.rect(0, 0, W, 6, stroke=0, fill=1)
    c.showPage(); c.save()
    return path

if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    mode = sys.argv[1] if len(sys.argv) > 1 else "all"
    if mode in ("all", "clean"):
        build_planner("en", os.path.join(OUT, "PlanaryShop-Budget-Planner-2026.xlsx"))
        build_planner("es", os.path.join(OUT, "PlanaryShop-Budget-Planner-2026-ES.xlsx"))
        build_guide("en", os.path.join(OUT, "Quick-Start-Guide.pdf"))
        build_guide("es", os.path.join(OUT, "Quick-Start-Guide-ES.pdf"))
        print("built clean EN+ES planners and guides")
    if mode == "sample":
        build_planner("en", sys.argv[2], sample=True)
        print("built sample-data planner:", sys.argv[2])
    if mode == "qa":
        # QA copy: inject 1000 income (Jan, source 1) + 300 expense (Jan, cat 1)
        p = sys.argv[2]
        build_planner("en", p)
        from openpyxl import load_workbook
        wb = load_workbook(p)
        wb["Income"]["C8"] = 1000
        wb["Expenses"]["C8"] = 300
        wb.save(p)
        print("built QA planner with test values:", p)
