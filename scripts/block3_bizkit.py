"""PlanaryShop — Block 3: Small Business Starter Kit (4 docx + Client Tracker xlsx + zip)."""
import os, zipfile
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SAGE = RGBColor(0x8A, 0x9B, 0x85); TERRA = RGBColor(0xC9, 0x8A, 0x6B)
INK = RGBColor(0x2E, 0x2A, 0x26); WGRAY = RGBColor(0x4A, 0x44, 0x3E)
BEIGE_HEX = "E9DFD2"; CREAM_HEX = "FAF7F2"
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "output", "producto-3")

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def base_doc():
    d = Document()
    st = d.styles['Normal']
    st.font.name = 'Lato'; st.font.size = Pt(10); st.font.color.rgb = INK
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(0.8)
        s.left_margin = s.right_margin = Inches(0.9)
    return d

def brand_header(d, doc_title):
    p = d.add_paragraph()
    r = p.add_run("PLANARYSHOP")
    r.font.name = 'Lato'; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = TERRA
    t = d.add_paragraph()
    r = t.add_run(doc_title)
    r.font.name = 'Georgia'; r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = INK
    rule = d.add_paragraph()
    rr = rule.add_run("─" * 58)
    rr.font.size = Pt(8); rr.font.color.rgb = SAGE
    rule.paragraph_format.space_after = Pt(10)

def h2(d, text, color=SAGE):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text.upper())
    r.font.name = 'Lato'; r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = color
    return p

def body(d, text, italic=False, size=10):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Lato'; r.font.size = Pt(size); r.font.italic = italic
    r.font.color.rgb = WGRAY if italic else INK
    return p

def field_row(table, i, label, value):
    c0, c1 = table.rows[i].cells
    r = c0.paragraphs[0].add_run(label)
    r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = WGRAY; r.font.name = 'Lato'
    r = c1.paragraphs[0].add_run(value)
    r.font.size = Pt(10); r.font.name = 'Lato'; r.font.color.rgb = INK
    shade(c0, BEIGE_HEX)

def footer_note(d, extra=None):
    if extra:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        r = p.add_run(extra)
        r.font.size = Pt(8.5); r.font.italic = True; r.font.color.rgb = WGRAY; r.font.name = 'Lato'
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    r = p.add_run("Thank you for supporting PlanaryShop! For personal use only.")
    r.font.name = 'Georgia'; r.font.size = Pt(9); r.font.italic = True; r.font.color.rgb = TERRA

# ---------------- Invoice ----------------
def build_invoice(path):
    d = base_doc()
    brand_header(d, "Invoice")
    t = d.add_table(rows=5, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.columns[0].width = Inches(1.8); t.columns[1].width = Inches(4.8)
    field_row(t, 0, "FROM", "[YOUR BUSINESS NAME] · [ADDRESS] · [EMAIL] · [PHONE]")
    field_row(t, 1, "BILL TO", "[CLIENT NAME] · [CLIENT ADDRESS] · [CLIENT EMAIL]")
    field_row(t, 2, "INVOICE #", "[INVOICE NUMBER]")
    field_row(t, 3, "DATE", "[ISSUE DATE]")
    field_row(t, 4, "DUE", "[DUE DATE]")
    h2(d, "Services")
    items = d.add_table(rows=6, cols=4)
    items.style = 'Table Grid'
    widths = [Inches(3.2), Inches(1.0), Inches(1.2), Inches(1.2)]
    for j, htxt in enumerate(["DESCRIPTION", "QTY", "RATE", "AMOUNT"]):
        cell = items.rows[0].cells[j]
        r = cell.paragraphs[0].add_run(htxt)
        r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(255,255,255)
        r.font.name = 'Lato'
        shade(cell, "8A9B85")
    for i in range(1, 4):
        vals = [f"[SERVICE OR ITEM {i}]", "[QTY]", "[$0.00]", "[$0.00]"]
        for j, v in enumerate(vals):
            r = items.rows[i].cells[j].paragraphs[0].add_run(v)
            r.font.size = Pt(9.5); r.font.name = 'Lato'
    for i, lab in ((4, "SUBTOTAL"), (5, "TAX ([RATE]%)")):
        r = items.rows[i].cells[2].paragraphs[0].add_run(lab)
        r.font.bold = True; r.font.size = Pt(9); r.font.name = 'Lato'; r.font.color.rgb = WGRAY
        r = items.rows[i].cells[3].paragraphs[0].add_run("[$0.00]")
        r.font.size = Pt(9.5); r.font.name = 'Lato'
        shade(items.rows[i].cells[2], BEIGE_HEX); shade(items.rows[i].cells[3], BEIGE_HEX)
    tot = items.add_row()
    r = tot.cells[2].paragraphs[0].add_run("TOTAL DUE")
    r.font.bold = True; r.font.size = Pt(10); r.font.name = 'Lato'; r.font.color.rgb = RGBColor(255,255,255)
    r = tot.cells[3].paragraphs[0].add_run("[$0.00]")
    r.font.bold = True; r.font.size = Pt(10); r.font.name = 'Lato'; r.font.color.rgb = RGBColor(255,255,255)
    shade(tot.cells[2], "C98A6B"); shade(tot.cells[3], "C98A6B")
    h2(d, "Payment terms", TERRA)
    body(d, "Payment is due by [DUE DATE] via [PAYMENT METHODS — e.g., bank transfer, PayPal, card]. "
            "Please reference invoice [INVOICE NUMBER] with your payment.")
    body(d, "Late payments may incur a fee of [LATE FEE]% per [WEEK/MONTH] past the due date.", italic=True)
    footer_note(d)
    d.save(path)

# ---------------- Quote ----------------
def build_quote(path):
    d = base_doc()
    brand_header(d, "Quote")
    t = d.add_table(rows=4, cols=2)
    field_row(t, 0, "PREPARED BY", "[YOUR BUSINESS NAME] · [EMAIL] · [PHONE]")
    field_row(t, 1, "PREPARED FOR", "[CLIENT NAME] · [CLIENT EMAIL]")
    field_row(t, 2, "QUOTE #", "[QUOTE NUMBER] — [DATE]")
    field_row(t, 3, "VALID UNTIL", "[EXPIRATION DATE] (30 days from issue unless noted)")
    h2(d, "Proposed services & pricing")
    items = d.add_table(rows=5, cols=2); items.style = 'Table Grid'
    for j, htxt in enumerate(["SERVICE", "PRICE"]):
        cell = items.rows[0].cells[j]
        r = cell.paragraphs[0].add_run(htxt)
        r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(255,255,255); r.font.name = 'Lato'
        shade(cell, "8A9B85")
    for i in range(1, 4):
        r = items.rows[i].cells[0].paragraphs[0].add_run(f"[SERVICE {i} — SHORT DESCRIPTION]")
        r.font.size = Pt(9.5); r.font.name = 'Lato'
        r = items.rows[i].cells[1].paragraphs[0].add_run("[$0.00]")
        r.font.size = Pt(9.5); r.font.name = 'Lato'
    r = items.rows[4].cells[0].paragraphs[0].add_run("ESTIMATED TOTAL")
    r.font.bold = True; r.font.size = Pt(10); r.font.name = 'Lato'; r.font.color.rgb = RGBColor(255,255,255)
    r = items.rows[4].cells[1].paragraphs[0].add_run("[$0.00]")
    r.font.bold = True; r.font.size = Pt(10); r.font.name = 'Lato'; r.font.color.rgb = RGBColor(255,255,255)
    shade(items.rows[4].cells[0], "C98A6B"); shade(items.rows[4].cells[1], "C98A6B")
    h2(d, "This quote includes")
    for it in ["[DELIVERABLE 1 — e.g., two design concepts]",
               "[DELIVERABLE 2 — e.g., one round of revisions]",
               "[DELIVERABLE 3 — e.g., final files in agreed formats]"]:
        p = d.add_paragraph(style='List Bullet')
        r = p.add_run(it); r.font.size = Pt(9.5); r.font.name = 'Lato'
    h2(d, "This quote does not include", TERRA)
    for it in ["[EXCLUSION 1 — e.g., additional revision rounds, billed at $[RATE]/hour]",
               "[EXCLUSION 2 — e.g., third-party costs such as printing, fonts, stock images]",
               "[EXCLUSION 3 — e.g., rush delivery]"]:
        p = d.add_paragraph(style='List Bullet')
        r = p.add_run(it); r.font.size = Pt(9.5); r.font.name = 'Lato'
    h2(d, "Acceptance")
    body(d, "By signing below, you accept this quote and authorize [YOUR BUSINESS NAME] to begin work "
            "under the terms above. A deposit of [DEPOSIT AMOUNT OR %] may be required before work begins.")
    sig = d.add_table(rows=2, cols=2)
    for j, lab in enumerate(["CLIENT SIGNATURE", "DATE"]):
        r = sig.rows[0].cells[j].paragraphs[0].add_run(lab)
        r.font.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = WGRAY; r.font.name = 'Lato'
        r = sig.rows[1].cells[j].paragraphs[0].add_run("_______________________________")
        r.font.color.rgb = WGRAY
    footer_note(d)
    d.save(path)

# ---------------- Service Agreement ----------------
def build_agreement(path):
    d = base_doc()
    brand_header(d, "Service Agreement")
    body(d, "This agreement is between [YOUR BUSINESS NAME] (“Provider”) and [CLIENT NAME] "
            "(“Client”), effective [START DATE].")
    clauses = [
        ("1. Services", "Provider will deliver the following services: [DESCRIBE SERVICES IN PLAIN WORDS]. "
         "Anything not listed here is out of scope and can be added by written agreement (email counts)."),
        ("2. Schedule", "Work begins on [START DATE] and is expected to finish by [END DATE]. "
         "Deadlines shift fairly if Client feedback or materials arrive late."),
        ("3. Payment", "Client agrees to pay [TOTAL PRICE] as follows: [PAYMENT SCHEDULE — e.g., 50% deposit, "
         "50% on delivery]. Invoices are due within [NUMBER] days. Work may pause on overdue balances."),
        ("4. Cancellations", "Either party may cancel with [NUMBER] days' written notice. Client pays for all "
         "work completed up to the cancellation date; deposits are [REFUNDABLE / NON-REFUNDABLE]."),
        ("5. Client responsibilities", "Client will provide the materials, access, and feedback Provider needs, "
         "within [NUMBER] business days of each request, so the project can stay on schedule."),
        ("6. Ownership", "Once paid in full, Client owns the final deliverables. Provider keeps ownership of "
         "drafts, unused concepts, and pre-existing tools, and may show the finished work in a portfolio "
         "unless Client asks otherwise in writing."),
        ("7. Liability", "Provider will do careful, professional work but is not liable for indirect losses "
         "(lost profits, lost data). Total liability is capped at the amount Client paid under this agreement."),
        ("8. Termination", "If either party seriously breaks this agreement and does not fix it within "
         "[NUMBER] days of written notice, the other party may end the agreement immediately."),
        ("9. Governing law", "This agreement is governed by the laws of [STATE], and any disputes will be "
         "handled in the courts of [COUNTY, STATE]."),
    ]
    for title, text in clauses:
        h2(d, title)
        body(d, text)
    sig = d.add_table(rows=2, cols=2)
    for j, lab in enumerate(["PROVIDER SIGNATURE / DATE", "CLIENT SIGNATURE / DATE"]):
        r = sig.rows[0].cells[j].paragraphs[0].add_run(lab)
        r.font.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = WGRAY; r.font.name = 'Lato'
        r = sig.rows[1].cells[j].paragraphs[0].add_run("_______________________________")
        r.font.color.rgb = WGRAY
    footer_note(d, "This is a template, not legal advice — have a local attorney review it before you rely on it.")
    d.save(path)

# ---------------- Cancellation Policy ----------------
def build_cancellation(path):
    d = base_doc()
    brand_header(d, "Cancellation Policy")
    body(d, "[YOUR BUSINESS NAME] · effective [DATE]. We know plans change — here is exactly how "
            "cancellations, reschedules, and refunds work, so there are never surprises.")
    h2(d, "Cancellation windows")
    t = d.add_table(rows=4, cols=2); t.style = 'Table Grid'
    rows = [("WHEN YOU CANCEL", "WHAT HAPPENS"),
            ("More than 48 hours before", "Full refund of your deposit, or free rescheduling — your choice."),
            ("24–48 hours before", "You may reschedule once for free. Cancelled appointments are charged "
             "[PERCENTAGE]% of the service price."),
            ("Less than 24 hours / no-show", "The full deposit is kept. No-shows are charged [PERCENTAGE]% "
             "of the service price.")]
    for i, (a, b) in enumerate(rows):
        ca, cb = t.rows[i].cells
        ra = ca.paragraphs[0].add_run(a); rb = cb.paragraphs[0].add_run(b)
        for rr in (ra, rb):
            rr.font.size = Pt(9.5); rr.font.name = 'Lato'
        if i == 0:
            ra.font.bold = rb.font.bold = True
            ra.font.color.rgb = rb.font.color.rgb = RGBColor(255,255,255)
            shade(ca, "8A9B85"); shade(cb, "8A9B85")
        else:
            ra.font.bold = True; shade(ca, BEIGE_HEX)
    h2(d, "Deposits")
    body(d, "A deposit of [DEPOSIT AMOUNT OR %] reserves your appointment. It counts toward your final "
            "price and is only kept in the situations listed above.")
    h2(d, "Rescheduling")
    body(d, "You can reschedule up to [NUMBER] times per booking through [BOOKING METHOD — e.g., our "
            "booking link or a quick message]. Reschedules requested with more than 24 hours' notice are free.")
    h2(d, "Refunds")
    body(d, "Approved refunds go back to the original payment method within [NUMBER] business days. "
            "If we ever need to cancel on you, you always receive a full refund or priority rebooking.")
    footer_note(d, "This is a template, not legal advice — have a local attorney review it before you rely on it.")
    d.save(path)

# ---------------- Client Tracker xlsx ----------------
def build_tracker(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    CREAM = "FAF7F2"; BEIGE = "E9DFD2"; SAGEH = "8A9B85"; TERRAH = "C98A6B"
    INKH = "2E2A26"; WGRAYH = "4A443E"; WHITE = "FFFFFF"
    thin = Side(style="thin", color=BEIGE)
    CUR = '$#,##0.00'
    wb = Workbook(); ws = wb.active; ws.title = "Clients"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = SAGEH
    for r in range(1, 48):
        for c in range(1, 14):
            ws.cell(row=r, column=c).fill = PatternFill("solid", start_color=CREAM)
    ws.column_dimensions['A'].width = 2
    t = ws.cell(row=2, column=2, value="CLIENT TRACKER")
    t.font = Font(name="Georgia", size=20, bold=True, color=INKH)
    s = ws.cell(row=3, column=2, value="One row per client — the Balance column and TOTALS row calculate themselves")
    s.font = Font(name="Lato", size=10, color=WGRAYH)
    for c in range(2, 13): ws.cell(row=4, column=c).fill = PatternFill("solid", start_color=SAGEH)
    ws.row_dimensions[4].height = 3
    heads = ["CLIENT","EMAIL","PHONE","SERVICE","STATUS","QUOTED","INVOICED","PAID",
             "BALANCE","NEXT FOLLOW-UP","NOTES"]
    widths = [18, 24, 14, 18, 12, 11, 11, 11, 11, 16, 26]
    from openpyxl.utils import get_column_letter
    for j, (h, w) in enumerate(zip(heads, widths)):
        col = 2 + j
        ws.column_dimensions[get_column_letter(col)].width = w
        cell = ws.cell(row=6, column=col, value=h)
        cell.font = Font(name="Lato", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", start_color=SAGEH)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[6].height = 20
    R0, RN = 7, 36  # 30 rows
    for r in range(R0, RN + 1):
        for j in range(len(heads)):
            col = 2 + j
            cell = ws.cell(row=r, column=col)
            if heads[j] == "BALANCE":
                cell.value = f"=H{r}-I{r}"
                cell.fill = PatternFill("solid", start_color=BEIGE)
            else:
                cell.fill = PatternFill("solid", start_color=WHITE)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            cell.font = Font(name="Lato", size=10, color=INKH)
            if heads[j] in ("QUOTED","INVOICED","PAID","BALANCE"):
                cell.number_format = CUR
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 18
    rt = RN + 1
    lab = ws.cell(row=rt, column=2, value="TOTALS")
    lab.font = Font(name="Lato", size=10, bold=True, color=WHITE)
    lab.fill = PatternFill("solid", start_color=WGRAYH)
    lab.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for colL in "GHIJ":
        cell = ws[f"{colL}{rt}"]
        cell.value = f"=SUM({colL}{R0}:{colL}{RN})"
        cell.number_format = CUR
        cell.font = Font(name="Lato", size=10, bold=True, color=INKH)
        cell.fill = PatternFill("solid", start_color=BEIGE)
        cell.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[rt].height = 20
    dv = DataValidation(type="list", formula1='"Lead,Quoted,Booked,In progress,Delivered,Paid"',
                        allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"F{R0}:F{RN}")
    ws.merge_cells(start_row=rt+2, start_column=2, end_row=rt+2, end_column=12)
    th = ws.cell(row=rt+2, column=2, value="Thank you for supporting PlanaryShop! For personal use only.")
    th.font = Font(name="Georgia", size=9, italic=True, color=TERRAH)
    th.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B7"
    wb.save(path)

if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    files = {
        "PlanaryShop-Invoice-Template.docx": build_invoice,
        "PlanaryShop-Quote-Template.docx": build_quote,
        "PlanaryShop-Service-Agreement-Template.docx": build_agreement,
        "PlanaryShop-Cancellation-Policy-Template.docx": build_cancellation,
        "PlanaryShop-Client-Tracker.xlsx": build_tracker,
    }
    for name, fn in files.items():
        fn(os.path.join(OUT, name))
        print("built", name)
    zpath = os.path.join(OUT, "PlanaryShop-Business-Kit.zip")
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for name in files:
            z.write(os.path.join(OUT, name), name)
    print("built PlanaryShop-Business-Kit.zip with", len(files), "files")
